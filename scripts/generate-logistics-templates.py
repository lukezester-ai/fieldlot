from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "public" / "documents"
OUTPUT.mkdir(parents=True, exist_ok=True)
FONT_PATH = ROOT / "public" / "fonts" / "NotoSans-Regular.ttf"
pdfmetrics.registerFont(TTFont("NotoSans", str(FONT_PATH)))

GREEN = colors.HexColor("#1B5E3C")
PALE = colors.HexColor("#EAF4EE")
GRID = colors.HexColor("#B8C4BC")
TEXT = colors.HexColor("#17221B")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_doc_font(document):
    styles = document.styles
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.size = Pt(20)
    styles["Heading 1"].font.size = Pt(14)
    styles["Heading 2"].font.size = Pt(11.5)


def add_field_table(document, rows):
    table = document.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(5.0)
    table.columns[1].width = Cm(11.5)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width = Cm(5.0)
        cells[1].width = Cm(11.5)
        cells[0].text = label
        cells[1].text = value
        cells[0].paragraphs[0].runs[0].bold = True
        set_cell_shading(cells[0], "EAF4EE")
        for cell in cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    document.add_paragraph()


def build_transport_agreement():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    set_doc_font(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Договор за автомобилен превоз на товари")
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.add_run("Примерна бланка за договаряне между възложител и превозвач").italic = True

    doc.add_heading("Страни по договора", level=1)
    add_field_table(doc, [
        ("Възложител", "[наименование, ЕИК, адрес, представител]"),
        ("Превозвач", "[наименование, ЕИК, лиценз, адрес, представител]"),
        ("Дата и място", "[дд.мм.гггг, населено място]"),
    ])

    doc.add_heading("Предмет и маршрут", level=1)
    add_field_table(doc, [
        ("Товар", "[вид, количество, опаковка, особености]"),
        ("Място на товарене", "[точен адрес и лице за контакт]"),
        ("Място на разтоварване", "[точен адрес и лице за контакт]"),
        ("Срок", "[дата и часови диапазон]"),
        ("Превозно средство", "[тип, регистрационен номер, товароносимост]"),
    ])

    clauses = [
        ("Цена и плащане", "Цената е [сума] лв. без/с ДДС и включва [пътни такси, чакане, товаро-разтоварни дейности]. Плащането се извършва по банков път в срок до [брой] дни след представяне на фактура и подписана товарителница."),
        ("Задължения на възложителя", "Възложителят осигурява вярна информация за товара, достъп до местата за товарене и разтоварване, необходимите документи и безопасни условия за работа."),
        ("Задължения на превозвача", "Превозвачът осигурява технически изправно и подходящо превозно средство, правоспособен водач, валидни лицензи и застраховки и спазва договорения маршрут и срок."),
        ("Отговорност и рекламации", "Страните носят отговорност съгласно приложимото законодателство и Конвенцията CMR, когато е приложима. Видими липси и повреди се отбелязват при приемането на товара."),
        ("Непреодолима сила", "Засегнатата страна уведомява другата страна без неоправдано забавяне и представя наличните доказателства за настъпилото обстоятелство."),
        ("Прекратяване и спорове", "Измененията се правят писмено. Споровете се решават чрез преговори, а при непостигане на съгласие - от компетентния български съд."),
    ]
    for heading, text in clauses:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(text)

    doc.add_heading("Подписи", level=1)
    signatures = doc.add_table(rows=2, cols=2)
    signatures.cell(0, 0).text = "За възложителя"
    signatures.cell(0, 1).text = "За превозвача"
    signatures.cell(1, 0).text = "Име: ____________________\nПодпис: __________________\nДата: ____________________"
    signatures.cell(1, 1).text = "Име: ____________________\nПодпис: __________________\nДата: ____________________"
    for cell in signatures.rows[0].cells:
        set_cell_shading(cell, "1B5E3C")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for row in signatures.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=140, bottom=140)

    note = doc.add_paragraph()
    note.add_run("Важно: ").bold = True
    note.add_run("Тази бланка е примерна и не представлява правен съвет. Адаптирайте я към конкретния превоз и я прегледайте с квалифициран юрист.")
    doc.save(OUTPUT / "transport-agreement-template.docx")


styles = getSampleStyleSheet()
TITLE = ParagraphStyle("TitleBG", parent=styles["Title"], fontName="NotoSans", fontSize=18, leading=22, textColor=TEXT, alignment=TA_CENTER, spaceAfter=8)
SUBTITLE = ParagraphStyle("SubBG", parent=styles["Normal"], fontName="NotoSans", fontSize=9, leading=12, textColor=colors.HexColor("#4E6255"), alignment=TA_CENTER, spaceAfter=10)
BODY = ParagraphStyle("BodyBG", parent=styles["Normal"], fontName="NotoSans", fontSize=8.5, leading=11, textColor=TEXT)
SMALL = ParagraphStyle("SmallBG", parent=BODY, fontSize=7.5, leading=9)
SECTION = ParagraphStyle("SectionBG", parent=BODY, fontSize=11, leading=14, textColor=GREEN, spaceBefore=8, spaceAfter=5)


def pdf_table(rows, widths, header=False, row_heights=None):
    table = Table(rows, colWidths=widths, rowHeights=row_heights, repeatRows=1 if header else 0)
    commands = [
        ("FONTNAME", (0, 0), (-1, -1), "NotoSans"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("TEXTCOLOR", (0, 0), (-1, -1), TEXT),
        ("GRID", (0, 0), (-1, -1), 0.5, GRID),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    if header:
        commands += [("BACKGROUND", (0, 0), (-1, 0), GREEN), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white)]
    table.setStyle(TableStyle(commands))
    return table


def footer(canvas, document):
    canvas.saveState()
    canvas.setFont("NotoSans", 7)
    canvas.setFillColor(colors.HexColor("#607066"))
    canvas.drawString(18 * mm, 10 * mm, "Fieldlot - примерна логистична бланка")
    canvas.drawRightString(192 * mm, 10 * mm, f"Страница {document.page}")
    canvas.restoreState()


def build_cmr():
    path = OUTPUT / "cmr-template.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=14 * mm, leftMargin=14 * mm, topMargin=14 * mm, bottomMargin=16 * mm)
    story = [Paragraph("CMR международна товарителница", TITLE), Paragraph("Работен образец за подготовка на данните по международен автомобилен превоз", SUBTITLE)]
    rows = [
        [Paragraph("1 Изпращач", BODY), Paragraph("2 Получател", BODY)],
        [Paragraph("Наименование и адрес:<br/><br/><br/>Държава:<br/>Лице за контакт:", BODY), Paragraph("Наименование и адрес:<br/><br/><br/>Държава:<br/>Лице за контакт:", BODY)],
        [Paragraph("3 Място на доставка", BODY), Paragraph("4 Място и дата на приемане", BODY)],
        [Paragraph("Адрес:<br/><br/>Държава:", BODY), Paragraph("Адрес:<br/><br/>Дата и час:", BODY)],
        [Paragraph("5 Приложени документи", BODY), Paragraph("6 Превозвач", BODY)],
        [Paragraph("Фактура / сертификат / опаковъчен лист:<br/><br/>", BODY), Paragraph("Наименование, адрес, лиценз:<br/><br/>", BODY)],
    ]
    story.append(pdf_table(rows, [91 * mm, 91 * mm], row_heights=[8 * mm, 24 * mm, 8 * mm, 16 * mm, 8 * mm, 17 * mm]))
    story += [Spacer(1, 5 * mm), Paragraph("Данни за товара", SECTION)]
    cargo = [[Paragraph(x, SMALL) for x in ["Марки и номера", "Брой и вид опаковки", "Описание на стоката", "Бруто кг", "Обем м3"]]]
    cargo += [["", "", "", "", ""] for _ in range(4)]
    story.append(pdf_table(cargo, [31 * mm, 35 * mm, 66 * mm, 25 * mm, 25 * mm], header=True, row_heights=[9 * mm] + [10 * mm] * 4))
    story += [Spacer(1, 5 * mm)]
    bottom = [
        [Paragraph("Инструкции на изпращача", BODY), Paragraph("Разходи и условия за плащане", BODY)],
        [Paragraph("<br/><br/><br/>", BODY), Paragraph("Навло:<br/>Допълнителни разходи:<br/>Платец:", BODY)],
        [Paragraph("Подпис и печат на изпращача<br/><br/><br/>", BODY), Paragraph("Подпис и печат на превозвача<br/><br/><br/>", BODY)],
        [Paragraph("Товарът е получен - дата, подпис и печат<br/><br/><br/>", BODY), Paragraph("Особени забележки и резерви<br/><br/><br/>", BODY)],
    ]
    story.append(pdf_table(bottom, [91 * mm, 91 * mm], row_heights=[8 * mm, 17 * mm, 17 * mm, 17 * mm]))
    story += [Spacer(1, 4 * mm), Paragraph("Този образец подпомага подготовката на данни и не замества официалните многолистни CMR формуляри или професионален правен съвет.", SMALL)]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def build_customs():
    path = OUTPUT / "customs-declaration-template.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=16 * mm, leftMargin=16 * mm, topMargin=15 * mm, bottomMargin=16 * mm)
    story = [Paragraph("Работен формуляр за митническа декларация", TITLE), Paragraph("Контролен лист за събиране на данни преди подаване чрез митнически представител или одобрена електронна система", SUBTITLE)]

    sections = [
        ("1 Участници", ["Износител / изпращач", "Вносител / получател", "EORI номера", "Митнически представител", "Лице за контакт"]),
        ("2 Пратка и транспорт", ["Референтен номер", "Държава на изпращане", "Държава на местоназначение", "Вид транспорт и регистрация", "Гранично митническо учреждение", "Incoterms и място"]),
        ("3 Стоки", ["Описание на стоката", "Код по Комбинираната номенклатура", "Произход", "Бруто и нето тегло", "Брой и вид опаковки", "Фактурна стойност и валута"]),
        ("4 Митнически режим", ["Заявен режим", "Предходен режим / документ", "Метод за митническа стойност", "Преференции и квоти", "Разрешения / сертификати"]),
    ]
    for index, (heading, labels) in enumerate(sections):
        if index == 3:
            story.append(PageBreak())
            story.append(Paragraph("Митнически данни и документи", TITLE))
        story.append(Paragraph(heading, SECTION))
        rows = [[Paragraph(label, BODY), ""] for label in labels]
        row_height = 9 * mm if index == 3 else 11 * mm
        story.append(pdf_table(rows, [62 * mm, 116 * mm], row_heights=[row_height] * len(rows)))
        story.append(Spacer(1, 3 * mm))

    story += [Paragraph("Приложени документи", SECTION)]
    docs = [
        "Търговска фактура", "Опаковъчен лист", "Транспортен документ / CMR", "Сертификат за произход",
        "Фитосанитарен или ветеринарен сертификат", "Лиценз / разрешение", "Застраховка", "Други",
    ]
    checks = [[Paragraph("☐", BODY), Paragraph(item, BODY), "Номер / дата: ______________________________"] for item in docs]
    story.append(pdf_table(checks, [10 * mm, 73 * mm, 95 * mm], row_heights=[9 * mm] * len(checks)))
    story += [Spacer(1, 3 * mm), Paragraph("Декларация на подготвилия данните", SECTION)]
    declaration = (
        "Потвърждавам, че въведените в този работен формуляр данни са проверени спрямо наличните търговски и транспортни документи. "
        "Разбирам, че официалната митническа декларация се подава единствено по приложимия законов ред и може да изисква допълнителни данни."
    )
    story += [Paragraph(declaration, BODY), Spacer(1, 4 * mm)]
    sign = [["Име и длъжност", "Дата", "Подпис"], ["", "", ""]]
    story.append(pdf_table(sign, [85 * mm, 40 * mm, 53 * mm], header=True, row_heights=[9 * mm, 18 * mm]))
    story += [Spacer(1, 3 * mm), Paragraph("Важно", SECTION), Paragraph("Това не е официален образец на митническа администрация и не може да бъде подадено като митническа декларация. Използвайте го само за подготовка и проверка на информацията с лицензиран митнически представител.", BODY)]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


if __name__ == "__main__":
    build_transport_agreement()
    build_cmr()
    build_customs()
    for path in sorted(OUTPUT.iterdir()):
        if path.name in {"transport-agreement-template.docx", "cmr-template.pdf", "customs-declaration-template.pdf"}:
            print(path)
